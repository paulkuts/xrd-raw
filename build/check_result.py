# -*- coding: utf-8 -*-
"""Проверка содержимого результатов идентификации: листы Excel, текст отчёта."""
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
sys.stdout.reconfigure(encoding="utf-8", errors="replace")

xlsx = ROOT / "input" / "sample1.xlsx"
wb = openpyxl.load_workbook(str(xlsx))
print("файл:", xlsx.name)
for name in wb.sheetnames:
    ws = wb[name]
    print("  лист '{}': строк {}, колонок {}".format(name, ws.max_row, ws.max_column))
    head = [str(c.value) for c in next(ws.iter_rows(min_row=1, max_row=1))]
    print("     шапка:", head[:8])
    if ws.max_row > 1:
        row = [str(c.value) for c in next(ws.iter_rows(min_row=2, max_row=2))]
        print("     строка 2:", row[:8])

txt = ROOT / "input" / "sample1_phases.txt"
raw = txt.read_bytes()
print("\nотчёт:", txt.name, len(raw), "байт, BOM:", raw[:3] == b"\xef\xbb\xbf")
text = raw.decode("utf-8-sig")
print("первые строки:")
for line in text.splitlines()[:6]:
    print("   ", line)
